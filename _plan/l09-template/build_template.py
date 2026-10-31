"""Builds the Lecture 9 practice set template (Verify Your Own Job) as an .xlsx.
Uploaded to Google Drive, it converts to a Google Sheet with the dropdowns intact.
Run: python build_template.py out.xlsx  (needs openpyxl)"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HEAD = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='2B2B2B')
EX = Font(italic=True, color='7A7A7A')
WRAP = Alignment(wrap_text=True, vertical='top')
ROWS = 80  # editable rows per table

def table(ws, head, widths, examples=(), dropdowns=None, rows=ROWS):
    ws.append(head)
    for c, w in zip(ws[1], widths):
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, WRAP
        ws.column_dimensions[c.column_letter].width = w
    for r in examples:
        ws.append(r)
        for c in ws[ws.max_row]:
            c.font = EX
    for row in ws.iter_rows(min_row=2, max_row=rows + 1, max_col=len(head)):
        for c in row:
            c.alignment = WRAP
    ws.freeze_panes = 'B2'
    for col, opts in (dropdowns or {}).items():
        dv = DataValidation(type='list', formula1='"' + ','.join(opts) + '"', allow_blank=True)
        dv.add(f'{col}2:{col}{rows + 1}')
        ws.add_data_validation(dv)

wb = Workbook()

# Read me
ws = wb.active
ws.title = 'Read me'
ws.column_dimensions['A'].width = 110
for line in [
    'Verify your own job · Lecture 9 practice set · 100xEngineers C8',
    'File > Make a copy. Grey italic rows are Aarav\'s worked example: overwrite or delete them.',
    'Brief and rules: https://100x-curriculum.vercel.app/c8/module-2/lecture-09-llms-from-first-principles/practice-set',
    '',
    '1. Numbers: write your prediction before anything else.',
    '2. Task sheet: one row per output from last week, from your records. Type, Cost if wrong, Data count and Bucket are dropdowns.',
    '3. A verifier that only answers in a quarter or longer is an outcome signal: type the row None.',
    '4. The other tabs are deliverables 3 to 9. Numbers fills itself from the task sheet.',
]:
    ws.append([line])
ws['A1'].font = Font(bold=True, size=14)

# Task sheet
ws = wb.create_sheet('Task sheet')
table(ws,
      ['Task', 'Minutes', 'Outcome', 'Verifier (judge · signal · when)', 'Type', 'Cost if wrong',
       'Data: where', 'Data: how many', 'Bucket', 'Notes (can a script produce it? flagged?)',
       "Judge's correction", 'Source (Track B)'],
      [34, 9, 26, 44, 11, 11, 30, 14, 11, 30, 26, 14],
      examples=[
          ["Pull last week's GMV, orders, AOV and return rate into the weekly business review deck", 90,
           'Leadership reviews correct numbers at Monday 11am',
           "Totals match the Databricks gold table; GMV within 0.5% of finance's SAP export",
           'Rule', 'Medium', 'Notebook history, 52 past weekly pulls', '10 to 100', 'Automate',
           'Scheduled notebook and sheet formula, no LLM'],
          ["Answer the category manager's Slack ask: why did Pune returns jump 38% last week?", 120,
           'Category gross margin',
           'Category manager · accepts the root cause and acts on it with no follow-up question · within 2 days',
           'Judgment', 'Medium', '#analytics-requests, each ask with his reply and what happened next',
           '100 to 1,000', 'Augment', 'Candidate task'],
          ['Pull Q3 sales by channel for the cost-cut deck', 60, 'Cost saved on the P&L',
           'Year-to-date minus Q1 and Q2 equals Q3; schema matches; within 0.5% of SAP',
           'Rule', 'High', 'Databricks, 11 past quarterly pulls', '10 to 100', 'Automate', 'No LLM'],
          ['Write SQL for repeat-purchase rate by acquisition channel for the growth team', 75,
           'CAC payback decisions',
           'Cohort sizes sum to total new customers; June cohort matches the signed-off Power BI number',
           'Rule', 'Medium', 'Team repo, about 60 saved queries with the question each answered', '10 to 100',
           'Automate', 'LLM drafts the SQL, the rule checks it'],
          ['Write user stories and acceptance criteria for "Buy 2, get 10% off" in Jira', 90,
           'Feature ships without rework',
           'Rule: every criterion is Given/When/Then. Judgment: PM and QA lead pass it at refinement with no clarifying questions · Thursday',
           'Mixed', 'Medium', 'Jira, about 200 past stories with reopen counts', '100 to 1,000', 'Augment', ''],
          ['Competitor pricing scan for the Q4 pricing recommendation', 180, 'Cost saved on the P&L',
           'Manager · would take it to the CFO, and every price has a dated source link · by Friday. Outcome signal: margin, Q3 to Q4',
           'Judgment', 'High', 'Drive, 6 past pricing docs with manager comments', 'Under 10', 'Augment', ''],
          ['A/B test readout for the one-page checkout', 120, 'Checkout conversion',
           'Rule: sample-ratio check passes, confidence interval computed correctly. Judgment: growth lead agrees with the ship call · same day',
           'Mixed', 'High', 'Confluence, 18 past readouts with the final decision', '10 to 100', 'Augment', ''],
          ['Build the Q4 cost-cut deck for leadership', 240, 'Cost saved on the P&L',
           'Rule: every number traces to the Q3 pull. Judgment: manager approves it for the CFO with no rework · same day',
           'Mixed', 'High', 'About 20 past leadership decks with version history', '10 to 100', 'Augment', ''],
          ['Present the recommendation to the CFO and handle pushback live', 60, 'Leadership decision',
           'None at task time. Outcome signal: savings realized next quarter',
           'None', 'High', '', 'None', 'Human', ''],
      ],
      dropdowns={'E': ['Rule', 'Judgment', 'Mixed', 'None'],
                 'F': ['Low', 'Medium', 'High'],
                 'H': ['None', 'Under 10', '10 to 100', '100 to 1,000', 'Over 1,000'],
                 'I': ['Automate', 'Augment', 'Human'],
                 'L': ['Interview', 'Artifact', 'Posting']})

# Numbers (deliverable 2), computed from the task sheet
T = "'Task sheet'!"
R = f'{ROWS + 1}'
ws = wb.create_sheet('Numbers')
ws.column_dimensions['A'].width = 52
ws.column_dimensions['B'].width = 40
verifiable = '+'.join(f'COUNTIFS({T}A2:A{R},"<>",{T}E2:E{R},"{t}")' for t in ('Rule', 'Judgment', 'Mixed'))
with_data = '+'.join(f'COUNTIFS({T}A2:A{R},"<>",{T}E2:E{R},"{t}",{T}H2:H{R},"<>None",{T}H2:H{R},"<>")'
                     for t in ('Rule', 'Judgment', 'Mixed'))
hours = '+'.join(f'SUMIFS({T}B2:B{R},{T}E2:E{R},"{t}")' for t in ('Rule', 'Judgment', 'Mixed'))
for row in [
    ['Deliverable 2 · your numbers', ''],
    ['Prediction: share of tasks verifiable (write first, never edit)', ''],
    ['Date of the prediction', ''],
    ['Meetings with no output of yours (hours)', ''],
    ['Tasks logged', f'=COUNTA({T}A2:A{R})'],
    ['Share of tasks verifiable', f'=IFERROR(({verifiable})/B5,"")'],
    ['Share of hours verifiable', f'=IFERROR(({hours})/SUM({T}B2:B{R}),"")'],
    ['Share verifiable, with data that exists', f'=IFERROR(({with_data})/B5,"")'],
    ['Which verdict surprised you, and why?', ''],
]:
    ws.append(row)
ws['A1'].font = Font(bold=True, size=13)
for c in ('B2', 'B6', 'B7', 'B8'):
    ws[c].number_format = '0%'
for r in range(2, 10):
    ws[f'A{r}'].alignment = WRAP

# Deliverable 3
ws = wb.create_sheet('Automated row')
table(ws, ['Automate row (from the task sheet)', 'How it runs without you', 'Minutes per week before',
           'Minutes per week after', 'Minutes saved'], [40, 44, 14, 14, 14],
      examples=[["Weekly business review pull", 'Scheduled notebook writes to the sheet; formula checks GMV against SAP',
                 90, 5, '=C2-D2']], rows=5)

# Deliverable 4
ws = wb.create_sheet('Candidate task')
ws.column_dimensions['A'].width = 40
ws.column_dimensions['B'].width = 80
ws.append(['Deliverable 4 · candidate task', ''])
ws['A1'].font = Font(bold=True, size=13)
for q in ['Task', 'Why it needs an LLM to produce', 'Verifier (judge · signal · when)', 'Where the data lives',
          'How many past examples', 'Confirmed I may use it on my own machine (who, when)',
          'Why it matters to my North Star']:
    ws.append([q, ''])
for r in range(2, 9):
    ws[f'A{r}'].alignment = ws[f'B{r}'].alignment = WRAP

# Deliverable 5
ws = wb.create_sheet('Checklists')
table(ws, ['Task', 'Line 1', 'Line 2', 'Line 3', 'Line 4', 'Line 5'], [34, 24, 24, 24, 24, 24],
      examples=[['"Why did metric X move?" Slack asks (candidate)', 'Driver named', 'Every number traced',
                 'Cause plausible', 'Action suggested', 'Under 150 words']], rows=5)

# Deliverable 6
ws = wb.create_sheet('Line test')
table(ws, ['#', 'Step', 'Orange or purple', 'If purple: why a rule cannot do it'], [5, 50, 16, 60],
      examples=[[1, 'Turn the ask into metric, region, period and comparison', 'Orange', ''],
                [2, 'Pull the metric and its baseline with SQL', 'Orange', ''],
                [3, 'Break the change down by SKU, courier, pincode and channel', 'Orange', ''],
                [4, 'Pick the driver that explains it and name the likely cause', 'Purple',
                 '"Because the size chart changed" comes from patterns in past cases'],
                [5, 'Write the reply for the category manager', 'Purple',
                 'What to say depends on who reads it and what they can act on'],
                [6, 'Check every number in the reply against step 3', 'Orange', '']],
      dropdowns={'C': ['Orange', 'Purple']}, rows=15)

# Deliverable 7
ws = wb.create_sheet('Golden pairs')
table(ws, ['#', 'Input', 'Ideal output', "Judge's verdict", 'Date of the case'], [5, 40, 56, 44, 14],
      examples=[[0, '"Pune returns up 38% week on week (week of 19 Oct). Why?" plus the step 3 breakdown',
                 '62% of the extra returns come from one SKU, women\'s straight-fit jeans in sizes 28 and 30, reason code "too small". Its size chart changed on 14 Oct. Couriers show no single driver. Suggest reverting the chart and rechecking next Monday.',
                 'Category manager reverted the chart on 21 Oct; returns were back to baseline within a week. Accepted, no follow-up.',
                 '19 Oct']] + [[i] for i in range(1, 11)], rows=12)

# Deliverable 8
ws = wb.create_sheet('My 20%')
table(ws, ['Judgment row (high cost if wrong, slow verifier)', 'Why it stays mine'], [56, 56], rows=8)
ws.append([])
ws.append(['My move for next month (own it, shadow it, or move the hours from deliverable 3)', ''])
ws.append(['By when', ''])

# Deliverable 9
ws = wb.create_sheet('Peer break')
table(ws, ['Peer', 'Verifier they tried to break', 'Their comment', 'My fix or defence'], [18, 40, 44, 44], rows=6)

# Stretch (section 07)
ws = wb.create_sheet('Baseline (optional)')
table(ws, ['Golden pair #', 'Template right?', 'Data right?', 'What it got wrong, and what it was missing'],
      [14, 14, 14, 70], dropdowns={'B': ['Yes', 'No'], 'C': ['Yes', 'No']}, rows=3)

wb.save(sys.argv[1])
